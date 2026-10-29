"""Independently verify the localised tables, workbook and figure inventory."""
from pathlib import Path
import argparse
import csv
import hashlib
import json
import math
import re
from datetime import datetime,timezone

from localisation import TABLES,NUMERIC,CJK,translate,source_tables_dir

def digest(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def check_metric_identities(rows):
    """Check identities from the supplied local results, without empirical constants."""
    assert rows, 'No model-strategy rows supplied'
    keys=[]
    sizes=[]
    for row in rows:
        key=row['Model-strategy key']
        keys.append(key)
        counts=[int(row[k]) for k in ['TN','FP','FN','TP']]
        assert all(value >= 0 for value in counts), (key, 'negative confusion count')
        tn,fp,fn,tp=counts
        n=sum(counts)
        assert n > 0, (key, 'empty evaluation set')
        sizes.append(n)
        for field in ['ROC-AUC','Average precision','Threshold']:
            value=float(row[field])
            assert math.isfinite(value) and 0 <= value <= 1, (key, field)
        for field,numerator,denominator in [('Sensitivity',tp,tp+fn),
                                            ('Specificity',tn,tn+fp),
                                            ('Records flagged per 1,000 records',1000*(tp+fp),n)]:
            if denominator:
                assert math.isclose(float(row[field]),numerator/denominator,
                                    rel_tol=1e-12,abs_tol=1e-11), (key, field)
    assert len(keys)==len(set(keys)), 'Duplicate model-strategy keys'
    assert len(set(sizes))==1, 'Models were evaluated on different-sized test sets'
    return len(rows)


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--source-results',required=True)
    parser.add_argument('--output',required=True)
    parser.add_argument('--report',type=Path,default=None,help='Optionally save the report; validation is otherwise read-only')
    args=parser.parse_args();source=Path(args.source_results);out=Path(args.output)
    rows_checked=0;numeric_checked=0;text_checked=0;payload=[]
    for zh,(en,sheet) in TABLES.items():
        with (source_tables_dir(source)/(zh+'.csv')).open(encoding='utf-8-sig',newline='') as fh:src=list(csv.reader(fh))
        with (out/'tables'/(en+'.csv')).open(encoding='utf-8-sig',newline='') as fh:dst=list(csv.reader(fh))
        assert len(src)==len(dst),(en,'row count')
        for sr,dr in zip(src,dst):
            assert len(sr)==len(dr)
            for s,d in zip(sr,dr):
                if NUMERIC.fullmatch(s) or s=='':assert s==d,(en,s,d);numeric_checked+=1
                else:assert translate(s)==d,(en,s,d);text_checked+=1
                assert not CJK.search(d),(en,d)
            rows_checked+=1
        values=[[c if i==0 or (c and not NUMERIC.fullmatch(c)) else None if c=='' else int(c) if re.fullmatch(r'[+-]?\d+',c) else float(c) for c in row] for i,row in enumerate(dst)]
        payload.append({'name':sheet,'values':values})
    with (out/'tables'/'main_test_metrics.csv').open(encoding='utf-8-sig',newline='') as fh:metrics=list(csv.DictReader(fh))
    metric_rows_checked=check_metric_identities(metrics)
    expected=['Figure_1_ROC','Figure_2_PR','Figure_3_calibration','Figure_4_SHAP_beeswarm','Figure_5_SHAP_waterfalls','Figure_6_SHAP_stability']
    for name in expected:
        for ext in ['.png','.pdf']:
            p=out/'figures'/'manuscript'/(name+ext);assert p.is_file() and p.stat().st_size>1000
    for ext in ['png','pdf']:assert len(list((out/'figures').glob('Analysis_*.'+ext)))==11
    text_audit=out/'reproducibility'/'figure_text_audit.json'
    label_scan_passed=None
    if text_audit.is_file():
        texts=json.loads(text_audit.read_text(encoding='utf-8'))
        assert not any(CJK.search(t) for labels in texts.values() for t in labels)
        label_scan_passed=True
    workbook_cells=None
    workbook=out/'analysis_results_English.xlsx'
    if workbook.exists():
        import openpyxl
        w=openpyxl.load_workbook(workbook,read_only=True,data_only=True)
        assert w.sheetnames==[x['name'] for x in payload]
        workbook_cells=0
        for item in payload:
            actual=list(w[item['name']].iter_rows(values_only=True))
            assert len(actual)==len(item['values'])
            for sr,dr in zip(item['values'],actual):
                assert len(sr)==len(dr)
                for s,d in zip(sr,dr):
                    if isinstance(s,(int,float)):
                        assert isinstance(d,(int,float)) and math.isclose(s,d,rel_tol=1e-12,abs_tol=1e-14),(item['name'],s,d)
                    else:assert s==d,(item['name'],s,d)
                    assert not isinstance(d,str) or not CJK.search(d)
                    workbook_cells+=1
        w.close()
    cache=source/'私密记录_禁止公开'/'绘图固定输入.joblib'
    report={'timestamp_utc':datetime.now(timezone.utc).isoformat(),'passed':True,'aggregate_tables':len(TABLES),
        'rows_checked_including_headers':rows_checked,'numeric_or_empty_csv_cells_unchanged':numeric_checked,
        'translated_text_cells_checked':text_checked,'workbook_cells_checked':workbook_cells,
        'model_strategy_rows_checked':metric_rows_checked,'metric_identities_checked':True,
        'manuscript_figures':6,'analysis_figures':11,'english_figure_label_scan_passed':label_scan_passed,
        'english_figure_text_audit_available':text_audit.is_file(),
        'model_training_rerun':False,'bootstrap_rerun':False,
        'private_plot_cache_available':cache.is_file(),
        'frozen_plot_cache_sha256':digest(cache) if cache.is_file() else None}
    if args.report:
        args.report.parent.mkdir(parents=True,exist_ok=True)
        args.report.write_text(json.dumps(report,indent=2),encoding='utf-8')
    print(json.dumps(report))

if __name__=='__main__':main()
