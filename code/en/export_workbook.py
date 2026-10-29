"""Portable value-only workbook export for the aggregate table JSON payload.

This fallback needs only openpyxl. It does not fit models, recompute statistics
or create figure previews. Existing published workbooks are not its inputs.
"""
from pathlib import Path
import argparse
import json
import math
import re


def export_workbook(output, input_file=None, filename='analysis_results_English.xlsx',
                    record_dir='reproducibility'):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    output = Path(output)
    records = output / record_dir
    input_file = Path(input_file) if input_file else records / 'workbook_input.json'
    payload = json.loads(input_file.read_text(encoding='utf-8'))
    if not isinstance(payload, list) or not payload:
        raise ValueError('The workbook payload must contain at least one table')
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = 'Research project'
    summaries = []
    for item in payload:
        values = item['values']
        if not values or not values[0] or any(len(r) != len(values[0]) for r in values):
            raise ValueError(f"Empty or non-rectangular table: {item['name']}")
        if item['name'] in workbook.sheetnames:
            raise ValueError(f"Duplicate worksheet name: {item['name']}")
        sheet = workbook.create_sheet(item['name'])
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = 'A2'
        for i, row in enumerate(values, 1):
            sheet.row_dimensions[i].height = 84 if i == 1 else 30
            for j, value in enumerate(row, 1):
                if isinstance(value, float) and not math.isfinite(value):
                    raise ValueError(f'Non-finite value in {item["name"]}!{i},{j}')
                cell = sheet.cell(i, j, value)
                # Aggregate labels are literal text, including an initial '='.
                if isinstance(value, str):
                    cell.data_type = 's'
                cell.font = Font(name='Arial', size=10, bold=i == 1,
                                 color='FFFFFF' if i == 1 else '172B4D')
                cell.alignment = Alignment(vertical='center', wrap_text=True,
                                           horizontal='center' if i == 1 else None)
                if i == 1:
                    cell.fill = PatternFill('solid', fgColor='17324D')
                elif i % 2 == 1:
                    cell.fill = PatternFill('solid', fgColor='F2F5F8')
        for j, label in enumerate(values[0], 1):
            column = [r[j-1] for r in values[1:]]
            numbers = [v for v in column if isinstance(v, (int, float)) and not isinstance(v, bool)]
            longest = max((len(v) for v in column if isinstance(v, str)), default=0)
            sheet.column_dimensions[get_column_letter(j)].width = 66 if longest > 90 else 42 if longest > 35 else 24
            if numbers:
                p_value = re.search(r'(?:^|\s)P value$', str(label)) or 'P值' in str(label)
                number_format = '0.00E+00' if p_value else '#,##0' if all(float(v).is_integer() for v in numbers) else '0.0000'
                for i in range(2, len(values)+1):
                    sheet.cell(i, j).number_format = number_format
                    sheet.cell(i, j).alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)
            for i, value in enumerate(column, 2):
                if isinstance(value, str) and len(value) > 90:
                    sheet.row_dimensions[i].height = max(sheet.row_dimensions[i].height, min(160, max(44, math.ceil(len(value)/85)*15)))
        summaries.append({'name': item['name'], 'rows': len(values)-1, 'columns': len(values[0])})
    output.mkdir(parents=True, exist_ok=True)
    records.mkdir(parents=True, exist_ok=True)
    destination = output / filename
    workbook.save(destination)
    # Reopen the exported values; formatting never changes the underlying data.
    saved = load_workbook(destination, read_only=True, data_only=False)
    checked = 0
    try:
        assert saved.sheetnames == [x['name'] for x in payload]
        for item in payload:
            actual = list(saved[item['name']].iter_rows(values_only=True))
            assert len(actual) == len(item['values'])
            for expected, row in zip(item['values'], actual):
                assert len(expected) == len(row)
                for a, b in zip(expected, row):
                    if isinstance(a, (int, float)) and not isinstance(a, bool):
                        assert isinstance(b, (int, float)) and math.isclose(a, b, rel_tol=1e-12, abs_tol=1e-14)
                    else:
                        assert (a if a != '' else None) == b
                    checked += 1
    finally:
        saved.close()
    report = {'backend': 'openpyxl', 'sheets': summaries, 'value_cells_checked': checked,
              'values_preserved': True, 'formula_cells_created': 0, 'layout_previews_generated': False}
    (records / 'workbook_export_checks.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('output', type=Path)
    parser.add_argument('--input', type=Path, default=None)
    parser.add_argument('--filename', default='analysis_results_English.xlsx')
    parser.add_argument('--record-dir', default='reproducibility')
    args = parser.parse_args()
    print(json.dumps(export_workbook(args.output, args.input, args.filename, args.record_dir), ensure_ascii=False))


if __name__ == '__main__':
    main()
